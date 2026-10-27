"""Turn uploaded files (PDF, DOCX, TXT/MD, images) into plain text.

Typed PDFs are read directly. Pages with little or no text layer (scans, handwriting)
and image files are rendered and transcribed by a vision model.
"""
from __future__ import annotations

import base64
from pathlib import Path
from typing import List, Optional

from .llm import LLMClient

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
TEXT_EXTS = {".txt", ".md", ".markdown"}
MIN_CHARS_PER_PAGE = 40  # below this, a PDF page is treated as scanned

TRANSCRIBE_SYSTEM = """You transcribe exam documents (often handwritten answer sheets) into plain text.
Rules:
- Transcribe EXACTLY what is written, including spelling mistakes and errors. Do not correct, complete or improve anything.
- Keep question numbers/labels exactly as they appear (e.g. "Q1", "2(b)") on their own lines.
- Keep the original line and paragraph order. Use LaTeX-style text for maths, e.g. x^2, sqrt(x), a/b.
- Write [illegible] for words you cannot read.
- For drawings/diagrams, write [diagram: short factual description of what is drawn and labelled].
- Ignore crossed-out text unless it is the only answer, then mark it [crossed out: ...].
Output only the transcription, no commentary."""


def _image_part(data: bytes, mime: str) -> dict:
    b64 = base64.b64encode(data).decode()
    return {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}}


def transcribe_images(llm: LLMClient, images: List[tuple], label: str) -> str:
    """images: list of (bytes, mime). Transcribes page by page to keep each call small."""
    pages = []
    for i, (data, mime) in enumerate(images, 1):
        content = [
            {"type": "text", "text": f"Transcribe page {i} of {len(images)} of the {label}."},
            _image_part(data, mime),
        ]
        text = llm.chat(llm.s.vision_model, TRANSCRIBE_SYSTEM, content,
                        task=f"transcribe:{label}", json_mode=False, max_tokens=4000)
        pages.append(text.strip())
    return "\n\n".join(pages)


def read_pdf(llm: Optional[LLMClient], data: bytes, label: str) -> str:
    import pymupdf as fitz

    doc = fitz.open(stream=data, filetype="pdf")
    out: List[str] = []
    for page in doc:
        text = page.get_text("text").strip()
        if len(text) >= MIN_CHARS_PER_PAGE or llm is None:
            out.append(text)
        else:
            png = page.get_pixmap(dpi=150).tobytes("png")
            out.append(transcribe_images(llm, [(png, "image/png")], label))
    return "\n\n".join(out).strip()


def read_docx(data: bytes) -> str:
    import io

    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts).strip()


def file_to_text(llm: Optional[LLMClient], name: str, data: bytes, label: str) -> str:
    ext = Path(name).suffix.lower()
    if ext in TEXT_EXTS:
        return data.decode("utf-8", errors="replace").strip()
    if ext == ".pdf":
        return read_pdf(llm, data, label)
    if ext == ".docx":
        return read_docx(data)
    if ext in IMAGE_EXTS:
        if llm is None:
            raise ValueError("Image input needs a vision model (LLM client).")
        mime = "image/jpeg" if ext in {".jpg", ".jpeg"} else f"image/{ext[1:]}"
        return transcribe_images(llm, [(data, mime)], label)
    raise ValueError(f"Unsupported file type for {label}: {name}")


def files_to_text(llm: Optional[LLMClient], files: List[tuple], label: str) -> str:
    """files: list of (name, bytes). Multiple files (e.g. one photo per page) are joined in order."""
    return "\n\n".join(file_to_text(llm, n, d, label) for n, d in files).strip()
